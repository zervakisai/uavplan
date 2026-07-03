#!/usr/bin/env python3
"""Run 2b: append verified references [44]-[59] (auto-numbered), add in-text
citations for the novelty claim (C3) and a UAS-sensing paragraph in §3.5 (C6)."""
from __future__ import annotations
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph

DOC = "revision/FLARE_Drones_REVISED.docx"

NEW_REFS = [
 "Boutilier, J.J.; Brooks, S.C.; Janmohamed, A.; Byers, A.; Buick, J.E.; Zhan, C.; Schoellig, A.P.; Cheskes, S.; Morrison, L.J.; Chan, T.C.Y. Optimizing a Drone Network to Deliver Automated External Defibrillators. Circulation 2017, 135, 2454–2465. https://doi.org/10.1161/CIRCULATIONAHA.116.026318",
 "Claesson, A.; Bäckman, A.; Ringh, M.; Svensson, L.; Nordberg, P.; Djärv, T.; Hollenberg, J. Time to Delivery of an Automated External Defibrillator Using a Drone for Simulated Out-of-Hospital Cardiac Arrests vs Emergency Medical Services. JAMA 2017, 317, 2332–2334. https://doi.org/10.1001/jama.2017.3957",
 "Nisingizwe, M.P.; et al. Effect of unmanned aerial vehicle (drone) delivery on blood product delivery time and wastage in Rwanda: a retrospective, cross-sectional study and time series analysis. Lancet Glob. Health 2022, 10, e564–e569. https://doi.org/10.1016/S2214-109X(22)00048-1",
 "Vansteenwegen, P.; Souffriau, W.; Van Oudheusden, D. The orienteering problem: A survey. Eur. J. Oper. Res. 2011, 209, 1–10. https://doi.org/10.1016/j.ejor.2010.03.045",
 "Gunawan, A.; Lau, H.C.; Vansteenwegen, P. Orienteering Problem: A survey of recent variants, solution approaches and applications. Eur. J. Oper. Res. 2016, 255, 315–332. https://doi.org/10.1016/j.ejor.2016.04.059",
 "Kelly, J.R.; Chaudhry, U.B. Mission-Driven UAV Path Selection: Post Hoc Cost Evaluation of Deterministic and Sampling Approaches. Drones 2026, 10, 152. https://doi.org/10.3390/drones10020152",
 "Blackmore, L.; Ono, M.; Williams, B.C. Chance-Constrained Optimal Path Planning With Obstacles. IEEE Trans. Robot. 2011, 27, 1080–1094. https://doi.org/10.1109/TRO.2011.2161160",
 "Rockafellar, R.T.; Uryasev, S. Optimization of Conditional Value-at-Risk. J. Risk 2000, 2, 21–41. https://doi.org/10.21314/JOR.2000.038",
 "Majumdar, A.; Pavone, M. How Should a Robot Assess Risk? Towards an Axiomatic Theory of Risk in Robotics. In Robotics Research (ISRR 2017); Springer: Cham, Switzerland, 2020; pp. 75–84. https://doi.org/10.1007/978-3-030-28619-4_10",
 "Karaman, S.; Frazzoli, E. Sampling-based algorithms for optimal motion planning. Int. J. Robot. Res. 2011, 30, 846–894. https://doi.org/10.1177/0278364911406761",
 "Jacobson, D.H. Optimal stochastic linear systems with exponential performance criteria and their relation to deterministic differential games. IEEE Trans. Autom. Control 1973, 18, 124–131. https://doi.org/10.1109/TAC.1973.1100265",
 "Whittle, P. Risk-sensitive linear/quadratic/Gaussian control. Adv. Appl. Probab. 1981, 13, 764–777. https://doi.org/10.2307/1426972",
 "Allison, R.S.; Johnston, J.M.; Craig, G.; Jennings, S. Airborne Optical and Thermal Remote Sensing for Wildfire Detection and Monitoring. Sensors 2016, 16, 1310. https://doi.org/10.3390/s16081310",
 "Thornberry, T.D.; et al. A Lightweight Remote Sensing Payload for Wildfire Detection and Fire Radiative Power Measurements. Sensors 2023, 23, 3514. https://doi.org/10.3390/s23073514",
 "Schuyler, T.J.; et al. Unmanned Aerial Systems for Monitoring Trace Tropospheric Gases. Atmosphere 2017, 8, 206. https://doi.org/10.3390/atmos8100206",
 "Guzman, M.I. Atmospheric Measurements with Unmanned Aerial Systems (UAS). Atmosphere 2020, 11, 1208. https://doi.org/10.3390/atmos11111208",
]


def replace_in_paragraph(p, old, new):
    full = "".join(r.text for r in p.runs)
    if old not in full:
        return False
    idx = full.index(old); end = idx + len(old); pos = 0; started = False
    for r in p.runs:
        rlen = len(r.text); rstart, rend = pos, pos + rlen; pos = rend
        if rend <= idx or rstart >= end:
            continue
        pre = r.text[:idx - rstart] if rstart < idx else ""
        suf = r.text[end - rstart:] if rend > end else ""
        r.text = (pre + new + suf) if not started else (pre + suf)
        started = True
    return True


def main():
    doc = Document(DOC)

    # locate last reference paragraph (auto-numbered list item = ref 43)
    last_ref = None
    for p in doc.paragraphs:
        if "San-Miguel-Ayanz, J.; et al. Forest Fires in Europe" in p.text:
            last_ref = p
    assert last_ref is not None, "ref 43 not found"

    # append new refs by deep-copying the numbered paragraph (keeps numPr/style)
    anchor_el = last_ref._p
    parent = last_ref._parent
    for reftext in NEW_REFS:
        newp = deepcopy(last_ref._p)
        anchor_el.addnext(newp)
        anchor_el = newp
        para = Paragraph(newp, parent)
        for r in list(para.runs):
            r.text = ""
        (para.runs[0] if para.runs else para.add_run("")).text = reftext
    print(f"[OK] appended {len(NEW_REFS)} references (now up to [{43+len(NEW_REFS)}])")

    # C3 in-text citation on the novelty sentence
    ok = False
    for p in doc.paragraphs:
        if replace_in_paragraph(p,
            "does evaluate mission-oriented outcomes;",
            "does evaluate mission-oriented outcomes [44–49];"):
            ok = True; break
    print("[OK] C3 novelty citation" if ok else "[MISS] C3 novelty citation")

    # C6 UAS-sensing sentences appended to §3.5 (freshness/ISR paragraph)
    ok = False
    for p in doc.paragraphs:
        if "reconnaissance (ISR) doctrine" in p.text and "situational awareness" in p.text:
            p.add_run(" In practice the achievable freshness is bounded by the sensor "
                "payload and its revisit time: onboard optical and thermal imagers, their "
                "swath and data latency, and payload size, weight and power limits govern how "
                "rapidly a fire perimeter can be re-surveyed [56–59]. The surveillance "
                "score applies no explicit fire-proximity weighting, a modeling simplification "
                "we note for transparency.")
            ok = True; break
    print("[OK] C6 sensing paragraph + [56-59]" if ok else "[MISS] C6 sensing")

    doc.save(DOC)
    print("\nSaved:", DOC)


if __name__ == "__main__":
    main()
